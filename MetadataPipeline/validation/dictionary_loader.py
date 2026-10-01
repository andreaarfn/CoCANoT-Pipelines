"""Load CoCANoT's machine-readable metadata dictionary from an XLSX workbook."""

from __future__ import annotations

from collections import OrderedDict
from pathlib import Path
from typing import Any, Iterable

import json
import openpyxl


MR_PREFIX = "MR "

DEFAULT_EXPECTED_TABLES = (
    "Clinical",
    "Surgical",
    "Electrophysiology",
    "Imaging",
)

JSON_COLUMNS = {
    "allowed_values_json": ("allowed_values", []),
    "conditional_allowed_values_json": ("conditional_allowed_values", {}),
    "exclusive_values_json": ("exclusive_values", []),
    "repeat_exclude_values_json": ("repeat_exclude_values", []),
    "repeat_special_values_json": ("repeat_special_values", []),
    "legacy_field_names_json": ("legacy_field_names", []),
}


class DictionaryLoadError(ValueError):
    """Raised when the XLSX cannot be safely converted into Python objects."""


def _clean_blank(value: Any) -> Any:
    if isinstance(value, str) and value.strip() == "":
        return None
    return value


def _parse_json_cell(
    value: Any,
    *,
    sheet_name: str,
    row_number: int,
    column_name: str,
    default: Any,
) -> Any:
    value = _clean_blank(value)

    if value is None:
        return default

    if isinstance(value, (list, dict)):
        return value

    if not isinstance(value, str):
        raise DictionaryLoadError(
            f"{sheet_name}! row {row_number}, column '{column_name}' must contain "
            f"JSON text or be blank; found {type(value).__name__}."
        )

    try:
        return json.loads(value)
    except json.JSONDecodeError as exc:
        raise DictionaryLoadError(
            f"Invalid JSON in {sheet_name}! row {row_number}, "
            f"column '{column_name}': {exc.msg}."
        ) from exc


def _read_headers(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[str]:
    raw_headers = [cell.value for cell in ws[1]]

    while raw_headers and raw_headers[-1] is None:
        raw_headers.pop()

    if not raw_headers:
        raise DictionaryLoadError(
            f"Machine-readable sheet '{ws.title}' has no headers."
        )

    headers: list[str] = []

    for column_number, value in enumerate(raw_headers, start=1):
        if value is None or str(value).strip() == "":
            raise DictionaryLoadError(
                f"Machine-readable sheet '{ws.title}' has a blank header "
                f"in column {column_number}."
            )
        headers.append(str(value).strip())

    duplicates = sorted(
        {
            header
            for header in headers
            if headers.count(header) > 1
        }
    )

    if duplicates:
        raise DictionaryLoadError(
            f"Machine-readable sheet '{ws.title}' has duplicate headers: "
            + ", ".join(duplicates)
        )

    return headers


def _read_machine_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
) -> dict[str, Any]:
    headers = _read_headers(ws)
    fields: list[dict[str, Any]] = []

    for row_number, row in enumerate(
        ws.iter_rows(
            min_row=2,
            max_col=len(headers),
            values_only=True,
        ),
        start=2,
    ):
        values = [_clean_blank(value) for value in row]

        if all(value is None for value in values):
            continue

        record = dict(zip(headers, values))

        for source_column, config in JSON_COLUMNS.items():
            output_key, default = config
            if source_column in record:
                record[output_key] = _parse_json_cell(
                    record[source_column],
                    sheet_name=ws.title,
                    row_number=row_number,
                    column_name=source_column,
                    default=default.copy(),
                )
            else:
                record[output_key] = default.copy()

        record["_sheet_name"] = ws.title
        record["_excel_row"] = row_number
        fields.append(record)

    return {
        "sheet_name": ws.title,
        "columns": headers,
        "fields": fields,
    }


def load_dictionary(
    workbook_path: str | Path,
    *,
    expected_tables: Iterable[str] = DEFAULT_EXPECTED_TABLES,
    reject_unexpected_mr_sheets: bool = True,
) -> dict[str, Any]:
    path = Path(workbook_path)

    if not path.exists():
        raise DictionaryLoadError(
            f"Dictionary file does not exist: {path}"
        )

    if path.suffix.lower() != ".xlsx":
        raise DictionaryLoadError(
            f"Dictionary must be an .xlsx workbook; received "
            f"'{path.suffix or 'no extension'}'."
        )

    expected_tables = tuple(expected_tables)
    expected_sheets = {
        f"{MR_PREFIX}{table}": table
        for table in expected_tables
    }

    try:
        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise DictionaryLoadError(
            f"Could not open dictionary workbook '{path}': {exc}"
        ) from exc

    try:
        machine_sheets = [
            name
            for name in workbook.sheetnames
            if name.startswith(MR_PREFIX)
        ]
        human_sheets = [
            name
            for name in workbook.sheetnames
            if not name.startswith(MR_PREFIX)
        ]

        missing = [
            sheet
            for sheet in expected_sheets
            if sheet not in machine_sheets
        ]

        if missing:
            raise DictionaryLoadError(
                "Dictionary is missing required machine-readable sheet(s): "
                + ", ".join(missing)
            )

        unexpected = [
            sheet
            for sheet in machine_sheets
            if sheet not in expected_sheets
        ]

        if unexpected and reject_unexpected_mr_sheets:
            raise DictionaryLoadError(
                "Dictionary contains unexpected machine-readable sheet(s): "
                + ", ".join(unexpected)
            )

        tables: OrderedDict[str, dict[str, Any]] = OrderedDict()

        for sheet_name, table_name in expected_sheets.items():
            tables[table_name] = _read_machine_sheet(
                workbook[sheet_name]
            )

        return {
            "source_file": str(path),
            "machine_readable_sheets": machine_sheets,
            "human_readable_sheets": human_sheets,
            "tables": tables,
        }
    finally:
        workbook.close()
